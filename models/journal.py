# models/journal.py — Journal des activités de la plateforme
import datetime
import uuid
from config import get_db

db = None
def _db():
    global db
    if db is None:
        db = get_db()
    return db

def _timestamp():
    return datetime.datetime.now().isoformat()

# Types d'actions enregistrées
ACTIONS = {
    "connexion":           "Connexion",
    "deconnexion":         "Déconnexion",
    "creation_compte":     "Création de compte",
    "suppression_compte":  "Suppression de compte",
    "suspension":          "Suspension de compte",
    "reactivation":        "Réactivation de compte",
    "creation_etab":       "Création d'établissement",
    "suspension_etab":     "Suspension d'établissement",
    "prolongement":        "Prolongement d'abonnement",
    "paiement":            "Paiement effectué",
    "exercice_publie":     "Exercice publié",
    "devoir_soumis":       "Devoir soumis",
    "correction":          "Correction enregistrée",
    "publication":         "Corrections publiées",
    "message_envoye":      "Message envoyé",
    "visio_lancee":        "Visioconférence lancée",
    "modification_forfait":"Modification des forfaits",
}

class Journal:
    """Enregistre toutes les activités de la plateforme."""

    @staticmethod
    def enregistrer(utilisateur_id: str, utilisateur_nom: str,
                    role: str, action: str, details: str = "",
                    ip: str = "") -> None:
        """Enregistre une action dans le journal."""
        entree = {
            "id":              str(uuid.uuid4()),
            "utilisateur_id":  utilisateur_id,
            "utilisateur_nom": utilisateur_nom,
            "role":            role,
            "action":          action,
            "libelle":         ACTIONS.get(action, action),
            "details":         details,
            "date":            _timestamp(),
            "ip":              ip,
        }
        try:
            _db().collection("journal").add(entree)
        except Exception:
            pass  # Ne jamais bloquer l'app à cause du journal

    @staticmethod
    def get_tout(limite: int = 500) -> list[dict]:
        """Récupère toutes les entrées du journal (super admin uniquement)."""
        try:
            docs = list(_db().collection("journal")
                        .order_by("date", direction="DESCENDING")
                        .limit(limite)
                        .stream())
            return [{"id": d.id, **d.to_dict()} for d in docs]
        except Exception:
            # Si index pas encore créé, récupérer sans tri
            docs = list(_db().collection("journal").limit(limite).stream())
            return sorted(
                [{"id": d.id, **d.to_dict()} for d in docs],
                key=lambda x: x.get("date", ""), reverse=True
            )

    @staticmethod
    def get_par_utilisateur(utilisateur_id: str) -> list[dict]:
        docs = list(_db().collection("journal")
                    .where("utilisateur_id", "==", utilisateur_id)
                    .stream())
        return sorted([{"id": d.id, **d.to_dict()} for d in docs],
                      key=lambda x: x.get("date", ""), reverse=True)

    @staticmethod
    def exporter_excel(entrees: list[dict]) -> bytes:
        """Exporte le journal en fichier Excel."""
        try:
            import io
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Journal des activités"

            # En-têtes
            entetes = ["Date", "Utilisateur", "Rôle", "Action", "Détails"]
            for col, entete in enumerate(entetes, 1):
                cell = ws.cell(row=1, column=col, value=entete)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="1A1A2E")
                cell.alignment = Alignment(horizontal="center")

            # Données
            for row, entree in enumerate(entrees, 2):
                ws.cell(row=row, column=1, value=entree.get("date", "")[:19].replace("T", " "))
                ws.cell(row=row, column=2, value=entree.get("utilisateur_nom", ""))
                ws.cell(row=row, column=3, value=entree.get("role", ""))
                ws.cell(row=row, column=4, value=entree.get("libelle", ""))
                ws.cell(row=row, column=5, value=entree.get("details", ""))

            # Largeurs colonnes
            ws.column_dimensions["A"].width = 20
            ws.column_dimensions["B"].width = 25
            ws.column_dimensions["C"].width = 20
            ws.column_dimensions["D"].width = 25
            ws.column_dimensions["E"].width = 40

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()
        except ImportError:
            # Si openpyxl pas installé, exporter en CSV
            import io
            output = io.StringIO()
            output.write("Date,Utilisateur,Rôle,Action,Détails\n")
            for e in entrees:
                date    = e.get("date","")[:19].replace("T"," ")
                nom     = e.get("utilisateur_nom","")
                role    = e.get("role","")
                action  = e.get("libelle","")
                details = e.get("details","").replace(",",";")
                output.write(f"{date},{nom},{role},{action},{details}\n")
            return output.getvalue().encode("utf-8")
